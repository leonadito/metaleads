"""The client's master RAG (Phase 7): documents about products, services,
prices and policies that the AI consults before answering a lead.

- Sources: a file (PDF, DOCX, TXT, MD), a web page, or pasted text. Only the
  extracted text is kept; the original file is not stored.
- The text is split into chunks of about CHUNK_CHARS (paragraph boundaries,
  with a small overlap); each chunk is embedded with the document's title in
  front, through the core's gateway (the agency's key, the client's quota).
- Search: the question is embedded with the same model and compared to the
  tenant's chunks by cosine distance in Postgres (pgvector). A tenant has few
  documents, so the search is exact and filtered by tenant, without an
  approximate index. Chunks indexed with another model are left out until
  the document is re-indexed.
- Web pages are fetched by the server, so only public addresses are allowed
  (no private, loopback or link-local hosts, checked on every redirect).
"""
import io
import ipaddress
import logging
import math
import re
import socket
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse

import requests
from django.db import connection, transaction
from django.utils import timezone
from django_tasks import task

from . import llm
from .models import KnowledgeChunk, KnowledgeDocument
from .vector import literal

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PAGE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 400_000
CHUNK_CHARS = 2000
OVERLAP_CHARS = 300
EMBED_BATCH = 64
TOP_K = 5
FILE_TYPES = {".pdf", ".docx", ".txt", ".md"}
MAX_REDIRECTS = 3


class RagError(Exception):
    pass


# ── Text extraction ──────────────────────────────────────────────────────────

def clean_text(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def extract_file(name, data):
    ext = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in FILE_TYPES:
        raise RagError("tipo de arquivo não aceito: use PDF, DOCX, TXT ou MD")
    if len(data) > MAX_FILE_BYTES:
        raise RagError("arquivo maior que 10 MB")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            import docx

            doc = docx.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except RagError:
        raise
    except Exception as exc:  # broken or protected files
        raise RagError(f"não consegui ler o arquivo: {exc}"[:300])
    text = clean_text(text)
    if len(text) < 20:
        raise RagError("o arquivo não tem texto legível (PDF digitalizado? envie a versão com texto)")
    return text


class _PageText(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "template", "head"}
    BLOCK = {"p", "div", "li", "br", "tr", "section", "article", "header", "footer", "h1", "h2", "h3", "h4", "h5",
             "h6", "table", "ul", "ol", "dd", "dt"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip, self.title, self.in_title = [], 0, "", False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self.in_title = True
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self.in_title = False
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip:
            self.parts.append(data)


def check_public_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise RagError("endereço inválido: use http:// ou https://")
    try:
        infos = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror:
        raise RagError("não encontrei esse endereço")
    for info in infos:
        if not ipaddress.ip_address(info[4][0].split("%")[0]).is_global:
            raise RagError("endereço não permitido (rede interna)")


def fetch_url(url):
    """(title, text) of a public web page (or PDF/text file)."""
    for _ in range(MAX_REDIRECTS + 1):
        check_public_url(url)
        try:
            response = requests.get(url, timeout=20, stream=True, allow_redirects=False,
                                    headers={"User-Agent": "AgitareCRM/1.0 (+https://agitare.com.br)"})
        except requests.RequestException as exc:
            raise RagError(f"não consegui abrir a página: {exc}"[:300])
        if response.is_redirect:
            url = urljoin(url, response.headers.get("Location", ""))
            continue
        break
    else:
        raise RagError("redirecionamentos demais")
    if response.status_code != 200:
        raise RagError(f"a página respondeu HTTP {response.status_code}")
    data = b""
    for block in response.iter_content(65536):
        data += block
        if len(data) > MAX_PAGE_BYTES:
            raise RagError("página maior que 5 MB")
    kind = (response.headers.get("Content-Type") or "").lower()
    name = urlparse(url).path.rsplit("/", 1)[-1] or "pagina"
    if "pdf" in kind:
        return name, extract_file(name if name.lower().endswith(".pdf") else name + ".pdf", data)
    if "text/plain" in kind:
        return name, extract_file(name + ".txt", data)
    parser = _PageText()
    parser.feed(data.decode(response.encoding or "utf-8", errors="replace"))
    text = clean_text("".join(parser.parts))
    if len(text) < 20:
        raise RagError("a página não tem texto legível (talvez seja montada por JavaScript)")
    return clean_text(parser.title)[:200], text


# ── Chunks ───────────────────────────────────────────────────────────────────

def split(text):
    """Chunks of about CHUNK_CHARS on paragraph boundaries, each starting with
    the end of the previous one (OVERLAP_CHARS) so a fact cut in two is not
    lost."""
    paragraphs = []
    for p in text.split("\n\n"):
        p = p.strip()
        while len(p) > CHUNK_CHARS:  # a huge paragraph: cut at a sentence or a space
            cut = max(p.rfind(". ", 0, CHUNK_CHARS), p.rfind(" ", 0, CHUNK_CHARS))
            cut = cut + 1 if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
            paragraphs.append(p[:cut].strip())
            p = p[cut:].strip()
        if p:
            paragraphs.append(p)
    chunks, current = [], ""
    for p in paragraphs:
        if current and len(current) + len(p) + 2 > CHUNK_CHARS:
            chunks.append(current)
            tail = current[-OVERLAP_CHARS:]
            current = tail[tail.find(" ") + 1:] if " " in tail else tail
        current = f"{current}\n\n{p}" if current else p
    if current:
        chunks.append(current)
    return chunks


def embed_all(tenant, texts):
    """Vectors for every text, in batches. Returns (vectors, model, tokens)."""
    vectors, model, tokens = [], "", 0
    for start in range(0, len(texts), EMBED_BATCH):
        answer = llm.embed(tenant, texts[start:start + EMBED_BATCH])
        vectors += answer["vectors"]
        model = answer.get("model", model)
        tokens += (answer.get("usage") or {}).get("total_tokens", 0)
    return vectors, model, tokens


@task
def index_document(document_id):
    doc = KnowledgeDocument.objects.select_related("tenant").filter(pk=document_id).first()
    if doc is None:
        return "sem documento"
    try:
        if doc.kind == KnowledgeDocument.Kind.URL and not doc.text:
            title, text = fetch_url(doc.source)
            doc.text = text
            doc.title = doc.title or title or doc.source[:200]
        pieces = split(doc.text)
        if not pieces:
            raise RagError("documento sem texto")
        vectors, model, tokens = embed_all(doc.tenant, [f"{doc.title}\n\n{p}" for p in pieces])
    except (RagError, llm.LLMError, requests.RequestException) as exc:
        KnowledgeDocument.objects.filter(pk=doc.pk).update(status=KnowledgeDocument.Status.ERROR,
                                                           error=str(exc)[:1000], title=doc.title or doc.source[:200])
        return f"erro: {exc}"
    with transaction.atomic():
        KnowledgeChunk.objects.filter(document=doc).delete()
        KnowledgeChunk.objects.bulk_create([
            KnowledgeChunk(document=doc, tenant=doc.tenant, order=i, text=p, embedding=v)
            for i, (p, v) in enumerate(zip(pieces, vectors))
        ])
        KnowledgeDocument.objects.filter(pk=doc.pk).update(
            title=doc.title, text=doc.text, status=KnowledgeDocument.Status.READY, error="", chunk_count=len(pieces),
            embedding_model=model, tokens=tokens, indexed_at=timezone.now())
    return f"{len(pieces)} trechos"


# ── Search ───────────────────────────────────────────────────────────────────

def _cosine(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    norm = math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b))
    return dot / norm if norm else 0.0


def search(tenant, query, k=TOP_K):
    """The k chunks closest to the query: [{"document", "text", "score"}].
    Returns {"results", "model", "stale"} where stale counts the ready
    documents indexed with another model (to re-index)."""
    ready = KnowledgeDocument.objects.filter(tenant=tenant, is_active=True, status=KnowledgeDocument.Status.READY)
    if not ready.exists() or not (query or "").strip():
        return {"results": [], "model": "", "stale": 0}
    answer = llm.embed(tenant, [query.strip()[:8000]])
    vector, model = answer["vectors"][0], answer.get("model", "")
    stale = ready.exclude(embedding_model=model).count()
    if connection.vendor == "postgresql":
        chunk, doc = KnowledgeChunk._meta.db_table, KnowledgeDocument._meta.db_table
        with connection.cursor() as cursor:
            cursor.execute(
                f"SELECT c.id, 1 - (c.embedding <=> %s::vector) FROM {chunk} c JOIN {doc} d ON d.id = c.document_id "
                f"WHERE c.tenant_id = %s AND d.is_active AND d.status = %s AND d.embedding_model = %s "
                f"AND c.embedding IS NOT NULL ORDER BY c.embedding <=> %s::vector LIMIT %s",
                [literal(vector), tenant.pk, KnowledgeDocument.Status.READY, model, literal(vector), k])
            scored = cursor.fetchall()
    else:  # SQLite in the tests
        rows = KnowledgeChunk.objects.filter(document__in=ready.filter(embedding_model=model)).exclude(embedding=None)
        scored = sorted(((c.pk, _cosine(vector, c.embedding)) for c in rows), key=lambda r: -r[1])[:k]
    chunks = KnowledgeChunk.objects.select_related("document").in_bulk([pk for pk, _ in scored])
    results = [{"document": chunks[pk].document.title, "text": chunks[pk].text, "score": round(float(score), 3)}
               for pk, score in scored if pk in chunks]
    return {"results": results, "model": model, "stale": stale}


def context_text(results):
    """The search results as the AI reads them, numbered for citation."""
    if not results:
        return "(nenhum trecho dos documentos do cliente)"
    return "\n\n".join(f"[documento {n}: {r['document']}]\n{r['text']}" for n, r in enumerate(results, start=1))
