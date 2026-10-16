import io
import socket
import uuid
from unittest import mock

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django_tasks_db.models import DBTaskResult

from leads.models import UserProfile

from . import llm, rag
from .models import KnowledgeChunk, KnowledgeDocument, Tenant


def fake_embed(model="openai/text-embedding-3-small"):
    """Deterministic 3-d vectors: texts about 'garagem' point one way, 'café' another."""
    def embed(tenant, texts, timeout=120):
        def vec(t):
            t = t.lower()
            return [1.0 if "garagem" in t or "vaga" in t else 0.0, 1.0 if "café" in t else 0.0, 0.1]
        return {"model": model, "vectors": [vec(t) for t in texts], "usage": {"total_tokens": 10 * len(texts)}}
    return embed


def public_dns(host, port, *args, **kwargs):
    return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", port))]


@override_settings(CORE_API_URL="https://core.example", CORE_WEBHOOK_SECRET="seg")
class RagTests(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(uuid=uuid.uuid4(), name="Alto", slug="alto")

    def test_split_keeps_paragraphs_with_overlap(self):
        text = "\n\n".join(f"Parágrafo {i}. " + "palavra " * 60 for i in range(12))
        chunks = rag.split(text)
        self.assertGreater(len(chunks), 3)
        self.assertTrue(all(len(c) <= rag.CHUNK_CHARS + rag.OVERLAP_CHARS for c in chunks))
        self.assertIn("palavra", chunks[1][:rag.OVERLAP_CHARS])  # starts with the end of the previous one
        self.assertEqual(rag.split("x" * 5000)[0], "x" * rag.CHUNK_CHARS)  # no spaces: hard cut

    def test_extract_txt_and_docx_and_reject_others(self):
        self.assertEqual(rag.extract_file("faq.txt", "Pergunta:   tem café?\r\n\r\n\r\nSim, das 7h às 10h.".encode()),
                         "Pergunta: tem café?\n\nSim, das 7h às 10h.")
        import docx

        d = docx.Document()
        d.add_paragraph("Apartamento de 3 quartos com vaga de garagem.")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Preço", "R$ 690 mil"
        buf = io.BytesIO()
        d.save(buf)
        text = rag.extract_file("tabela.docx", buf.getvalue())
        self.assertIn("vaga de garagem", text)
        self.assertIn("Preço | R$ 690 mil", text)
        with self.assertRaisesMessage(rag.RagError, "tipo de arquivo"):
            rag.extract_file("planilha.xlsx", b"x")
        with self.assertRaisesMessage(rag.RagError, "não tem texto"):
            rag.extract_file("vazio.txt", b"  ")

    def test_only_public_addresses_even_after_a_redirect(self):
        for address in ("127.0.0.1", "10.0.0.5", "169.254.169.254", "::1", "172.18.0.3"):
            with mock.patch("crm.rag.socket.getaddrinfo",
                            return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, 80))]):
                with self.assertRaisesMessage(rag.RagError, "rede interna"):
                    rag.check_public_url("http://site.example/")
        with self.assertRaisesMessage(rag.RagError, "http"):
            rag.check_public_url("file:///etc/passwd")
        redirect = mock.Mock(is_redirect=True, headers={"Location": "http://interno.example/admin"})

        def dns(host, port, *a, **k):
            ip = "10.0.0.9" if host == "interno.example" else "93.184.216.34"
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port))]
        with mock.patch("crm.rag.socket.getaddrinfo", side_effect=dns), \
                mock.patch("crm.rag.requests.get", return_value=redirect) as get:
            with self.assertRaisesMessage(rag.RagError, "rede interna"):
                rag.fetch_url("https://site.example/")
        self.assertEqual(get.call_count, 1)  # the internal address is never requested
        self.assertFalse(get.call_args.kwargs["allow_redirects"])

    def test_web_page_text(self):
        html = ("<html><head><title>Pousada Sol</title><style>x{}</style></head><body><nav>Menu</nav>"
                "<h1>Quartos</h1><p>Café da manhã incluso, das 7h às 10h.</p><script>alert(1)</script></body></html>")
        page = mock.Mock(is_redirect=False, status_code=200, headers={"Content-Type": "text/html"}, encoding="utf-8",
                         iter_content=lambda n: [html.encode()])
        with mock.patch("crm.rag.socket.getaddrinfo", side_effect=public_dns), \
                mock.patch("crm.rag.requests.get", return_value=page):
            title, text = rag.fetch_url("https://pousada.example/quartos")
        self.assertEqual(title, "Pousada Sol")
        self.assertIn("Café da manhã incluso", text)
        self.assertNotIn("alert", text)

    def test_index_and_search_by_meaning_with_the_same_model(self):
        garage = KnowledgeDocument.objects.create(tenant=self.tenant, kind="text", title="Estrutura",
                                                  text="O prédio tem vaga de garagem coberta para cada unidade.")
        breakfast = KnowledgeDocument.objects.create(tenant=self.tenant, kind="text", title="Serviços",
                                                     text="O café da manhã é servido das 7h às 10h.")
        with mock.patch.object(llm, "embed", side_effect=fake_embed()):
            self.assertEqual(rag.index_document.call(garage.pk), "1 trechos")
            rag.index_document.call(breakfast.pk)
            found = rag.search(self.tenant, "tem garagem?")
        garage.refresh_from_db()
        self.assertEqual((garage.status, garage.chunk_count, garage.tokens), ("ready", 1, 10))
        self.assertEqual(KnowledgeChunk.objects.get(document=garage).embedding, [1.0, 0.0, 0.1])
        self.assertEqual(found["results"][0]["document"], "Estrutura")
        self.assertEqual(found["stale"], 0)
        self.assertIn("[documento 1: Estrutura]", rag.context_text(found["results"]))
        # the agency switched the embedding model: old chunks leave the search until re-indexed
        with mock.patch.object(llm, "embed", side_effect=fake_embed("qwen/qwen3-embedding-0.6b")):
            found = rag.search(self.tenant, "tem garagem?")
        self.assertEqual((found["results"], found["stale"]), ([], 2))
        # inactive documents are not searched
        KnowledgeDocument.objects.filter(pk=garage.pk).update(is_active=False)
        with mock.patch.object(llm, "embed", side_effect=fake_embed()):
            self.assertEqual([r["document"] for r in rag.search(self.tenant, "garagem")["results"]], ["Serviços"])

    def test_index_error_is_shown_on_the_document(self):
        doc = KnowledgeDocument.objects.create(tenant=self.tenant, kind="text", title="X", text="texto qualquer aqui")
        with mock.patch.object(llm, "embed", side_effect=llm.QuotaExhausted("cota mensal esgotada")):
            rag.index_document.call(doc.pk)
        doc.refresh_from_db()
        self.assertEqual((doc.status, doc.error), ("error", "cota mensal esgotada"))


@override_settings(CORE_API_URL="https://core.example", CORE_WEBHOOK_SECRET="seg")
class RagScreenTests(TestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(uuid=uuid.uuid4(), name="Alto", slug="alto")
        user = User.objects.create_user(username="dona")
        UserProfile.objects.create(user=user, core_client_uuid=self.tenant.uuid)
        self.client.force_login(user)

    def test_upload_url_text_and_actions(self):
        f = SimpleUploadedFile("faq.txt", "Aceitamos pets de pequeno porte, com taxa de R$ 50.".encode())
        r = self.client.post("/crm/ia/documentos/novo/", {"kind": "file", "file": f})
        self.assertIn("ok=", r["Location"])
        doc = KnowledgeDocument.objects.get()
        self.assertEqual((doc.title, doc.kind, doc.created_by), ("faq", "file", "dona"))
        self.assertTrue(DBTaskResult.objects.filter(task_path="crm.rag.index_document").exists())
        with mock.patch("crm.rag.socket.getaddrinfo",
                        return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("192.168.0.1", 80))]):
            r = self.client.post("/crm/ia/documentos/novo/", {"kind": "url", "url": "http://roteador.example/"})
        self.assertIn("erro=", r["Location"])
        self.client.post("/crm/ia/documentos/novo/", {"kind": "text", "title": "Regras", "text": "Check-in às 14h e check-out às 12h."})
        self.assertEqual(KnowledgeDocument.objects.count(), 2)
        self.client.post(f"/crm/ia/documentos/{doc.uuid}/uso/")
        doc.refresh_from_db()
        self.assertFalse(doc.is_active)
        other = Tenant.objects.create(uuid=uuid.uuid4(), name="Outro", slug="outro")
        foreign = KnowledgeDocument.objects.create(tenant=other, kind="text", title="Alheio", text="x" * 30)
        self.assertEqual(self.client.post(f"/crm/ia/documentos/{foreign.uuid}/apagar/").status_code, 404)
        page = self.client.get("/crm/ia/documentos/")
        self.assertContains(page, "Regras")
        self.assertNotContains(page, "Alheio")
